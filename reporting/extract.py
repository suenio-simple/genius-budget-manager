"""
Budget Manager — Módulo de reportería
Extrae datos de la API REST y genera un archivo Excel con métricas de campañas.

Uso:
    python extract.py

Requisitos:
    pip install -r requirements.txt

El servidor de Budget Manager debe estar corriendo en http://localhost:8080
"""

import json
import sys
import urllib.request
import urllib.error
from datetime import datetime

BUDGET_MANAGER_URL = 'http://localhost:8080'
LANDING_CRM_URL    = 'http://localhost:3000'
OUTPUT_FILE        = 'report.xlsx'

REQUIRED_CAMPAIGN_FIELDS = ('id', 'name', 'client', 'status', 'budget', 'spent')
REQUIRED_SUMMARY_FIELDS  = ('activeCampaigns', 'totalBudget', 'totalSpent', 'totalAvailable', 'consumptionPercentage')


def api_get(url: str) -> list | dict:
    try:
        with urllib.request.urlopen(url, timeout=5) as response:
            return json.loads(response.read())
    except (urllib.error.URLError, TimeoutError) as e:
        raise RuntimeError(f'No se pudo conectar a {url}: {e}') from e
    except json.JSONDecodeError as e:
        raise RuntimeError(f'Respuesta inválida (no es JSON) desde {url}: {e}') from e


def get_campaigns() -> list:
    """Obtiene todas las campañas del Budget Manager."""
    return api_get(f'{BUDGET_MANAGER_URL}/api/campaigns')


def get_budget_summary() -> dict:
    """Obtiene el resumen global de presupuesto (solo campañas activas)."""
    return api_get(f'{BUDGET_MANAGER_URL}/api/campaigns/summary')


def validate_campaigns(campaigns) -> None:
    """Verifica que la lista de campañas tenga la forma que espera el reporte."""
    if not isinstance(campaigns, list):
        raise RuntimeError(f'Se esperaba una lista de campañas, se recibió: {type(campaigns).__name__}')
    for i, c in enumerate(campaigns):
        if not isinstance(c, dict):
            raise RuntimeError(f'Campaña #{i} no es un objeto válido: {c!r}')
        missing = [f for f in REQUIRED_CAMPAIGN_FIELDS if f not in c]
        if missing:
            raise RuntimeError(f'Campaña #{i} (id={c.get("id", "?")}) no tiene los campos: {", ".join(missing)}')
        for field in ('budget', 'spent'):
            if not isinstance(c[field], (int, float)):
                raise RuntimeError(f'Campaña #{i} (id={c.get("id")}) tiene "{field}" no numérico: {c[field]!r}')


def validate_summary(summary) -> None:
    """Verifica que el resumen de presupuesto tenga la forma que espera el reporte."""
    if not isinstance(summary, dict):
        raise RuntimeError(f'Se esperaba un objeto de resumen, se recibió: {type(summary).__name__}')
    missing = [f for f in REQUIRED_SUMMARY_FIELDS if f not in summary]
    if missing:
        raise RuntimeError(f'El resumen no tiene los campos: {", ".join(missing)}')


def get_leads_summary() -> list:
    """Obtiene el resumen de leads por landing desde el Landing CRM."""
    return api_get(f'{LANDING_CRM_URL}/api/landings/summary')


def _display_length(cell) -> int:
    """Longitud aproximada del contenido ya formateado (moneda/porcentaje)."""
    if isinstance(cell.value, (int, float)) and '%' in cell.number_format:
        return len(f'{cell.value:,.2f}%')
    if isinstance(cell.value, (int, float)) and '$' in cell.number_format:
        return len(f'$ {cell.value:,.2f}')
    return len(str(cell.value))


def autofit_columns(ws) -> None:
    """Ajusta el ancho de cada columna al contenido más largo, ya formateado."""
    for column_cells in ws.columns:
        length = max(_display_length(cell) for cell in column_cells if cell.value is not None)
        ws.column_dimensions[column_cells[0].column_letter].width = length + 2


def export_to_excel(campaigns: list, summary: dict) -> None:
    """Genera el archivo Excel con métricas de campañas."""
    import openpyxl
    from openpyxl.styles import Font, Border, Side
    wb = openpyxl.Workbook()
    bold = Font(bold=True)
    thin = Side(style='thin')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Hoja de campañas
    ws = wb.active
    ws.title = 'Campañas'
    ws.append(['ID', 'Nombre', 'Cliente', 'Estado', 'Presupuesto', 'Gastado', 'Disponible'])
    for cell in ws[1]:
        cell.font = bold
    for c in campaigns:
        ws.append([
            c.get('id'),
            c.get('name'),
            c.get('client'),
            c.get('status'),
            c.get('budget', 0),
            c.get('spent', 0),
            c.get('budget', 0) - c.get('spent', 0),
        ])

    money_format = '"$" #,##0.00'
    for row in ws.iter_rows(min_row=2, min_col=5, max_col=7, max_row=ws.max_row):
        for cell in row:
            cell.number_format = money_format

    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=7):
        for cell in row:
            cell.border = border

    # Hoja de resumen
    ws2 = wb.create_sheet('Resumen')
    ws2.append(['Métrica', 'Valor'])
    for cell in ws2[1]:
        cell.font = bold
    ws2.append(['Campañas activas',    summary.get('activeCampaigns', 0)])
    ws2.append(['Presupuesto total',   summary.get('totalBudget', 0)])
    ws2.append(['Total gastado',       summary.get('totalSpent', 0)])
    ws2.append(['Total disponible',    summary.get('totalAvailable', 0)])
    ws2.append(['% de consumo',        summary.get('consumptionPercentage', 0)])

    for row in ws2.iter_rows(min_row=3, max_row=5, min_col=2, max_col=2):
        for cell in row:
            cell.number_format = money_format

    ws2.cell(row=6, column=2).number_format = '0.00"%"'

    for row in ws2.iter_rows(min_row=1, max_row=ws2.max_row, min_col=1, max_col=2):
        for cell in row:
            cell.border = border

    autofit_columns(ws)
    autofit_columns(ws2)

    wb.save(OUTPUT_FILE)
    print(f'Reporte guardado en {OUTPUT_FILE}')


if __name__ == '__main__':
    try:
        print(f'Extrayendo datos — {datetime.now().strftime("%Y-%m-%d %H:%M")}')
        campaigns = get_campaigns()
        validate_campaigns(campaigns)
        summary = get_budget_summary()
        validate_summary(summary)
        print(f'Campañas encontradas: {len(campaigns)}')
        export_to_excel(campaigns, summary)
    except RuntimeError as e:
        print(f'Error: {e}', file=sys.stderr)
        sys.exit(1)
